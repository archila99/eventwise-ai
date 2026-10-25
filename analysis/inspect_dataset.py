"""Lightweight, read-only structural inspection of the PDF/DOCX files under data/.

Usage:  .venv/bin/python analysis/inspect_dataset.py [data_dir] [out_json]

Collects per-file metrics (text density, tables, images, vector drawings,
form fields, headers/footers, merged cells, checkboxes, charts) without
modifying the source documents.
"""

import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

import docx
import pymupdf

# ballot boxes, check marks and squares (escaped so the file stays ASCII-safe)
CHECKBOX_CHARS = "\u2610\u2611\u2612\u2713\u2714\u2717\u2718\u25a1\u25a0"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def inspect_pdf(path: Path) -> dict:
    doc = pymupdf.open(path)
    pages = []
    first_lines, last_lines = Counter(), Counter()
    for page in doc:
        text = page.get_text()
        lines = [l.strip() for l in text.splitlines() if l.strip()]
        if lines:
            first_lines[re.sub(r"\d+", "#", lines[0])] += 1
            last_lines[re.sub(r"\d+", "#", lines[-1])] += 1
        try:
            tables = page.find_tables().tables
        except Exception:
            tables = []
        drawings = page.get_drawings()
        pages.append({
            "chars": len(text.strip()),
            "images": len(page.get_images(full=True)),
            "drawings": len(drawings),
            "tables": [(t.row_count, t.col_count) for t in tables],
            "widgets": [(w.field_type_string, w.field_name) for w in page.widgets()],
            "checkbox_chars": sum(text.count(c) for c in CHECKBOX_CHARS),
        })
    n = len(pages)
    fonts = {f[3] for p in doc for f in p.get_fonts()}
    result = {
        "type": "pdf",
        "pages": n,
        "producer": doc.metadata.get("producer"),
        "creator": doc.metadata.get("creator"),
        "toc_entries": len(doc.get_toc()),
        "total_chars": sum(p["chars"] for p in pages),
        "low_text_pages": [i + 1 for i, p in enumerate(pages) if p["chars"] < 50],
        "pages_with_images": sum(1 for p in pages if p["images"]),
        "total_images": sum(p["images"] for p in pages),
        "pages_with_many_drawings": [i + 1 for i, p in enumerate(pages) if p["drawings"] > 40],
        "tables_found": sum(len(p["tables"]) for p in pages),
        "table_shapes": [s for p in pages for s in p["tables"]][:15],
        "widgets": [w for p in pages for w in p["widgets"]][:20],
        "checkbox_chars": sum(p["checkbox_chars"] for p in pages),
        "repeated_first_lines": [l for l, c in first_lines.items() if n > 2 and c >= n * 0.5][:3],
        "repeated_last_lines": [l for l, c in last_lines.items() if n > 2 and c >= n * 0.5][:3],
        "font_count": len(fonts),
    }
    doc.close()
    return result


def inspect_docx(path: Path) -> dict:
    d = docx.Document(path)
    styles = Counter(p.style.name for p in d.paragraphs if p.text.strip())
    headings = [p.text.strip()[:60] for p in d.paragraphs
                if p.style.name.lower().startswith(("heading", "title")) and p.text.strip()]
    tables = []
    merged = 0
    nested = 0
    for t in d.tables:
        xml = t._tbl.xml
        merged += xml.count("w:gridSpan") + xml.count("w:vMerge")
        nested += xml.count("<w:tbl>") + xml.count("<w:tbl ") - 1
        tables.append((len(t.rows), len(t.columns)))
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        body = z.read("word/document.xml").decode("utf8", "ignore")
        header_text = " | ".join(
            re.sub(r"<[^>]+>", "", z.read(n).decode("utf8", "ignore")).strip()[:80]
            for n in names if re.match(r"word/(header|footer)\d*\.xml", n)
        )
    full_text = "\n".join(p.text for p in d.paragraphs)
    table_text = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    langs = Counter(re.findall(r'w:lang w:val="([^"]+)"', body))
    return {
        "type": "docx",
        "paragraphs": sum(1 for p in d.paragraphs if p.text.strip()),
        "para_chars": len(full_text.strip()),
        "table_chars": len(table_text.strip()),
        "headings": len(headings),
        "heading_sample": headings[:6],
        "top_styles": styles.most_common(6),
        "tables": len(tables),
        "table_shapes": tables[:15],
        "merged_cell_markers": merged,
        "nested_tables": nested,
        "media_files": [n.split("/")[-1] for n in names if n.startswith("word/media/")],
        "charts": [n for n in names if n.startswith("word/charts/chart")],
        "inline_drawings": body.count("<w:drawing>"),
        "sdt_content_controls": body.count("<w:sdt>"),
        "w14_checkboxes": body.count("w14:checkbox"),
        "legacy_form_fields": body.count("w:fldChar") and body.count("FORMCHECKBOX"),
        "checkbox_chars": sum((full_text + table_text).count(c) for c in CHECKBOX_CHARS),
        "sections": len(d.sections),
        "header_footer_text": header_text[:200],
        "explicit_page_breaks": body.count('w:type="page"') + body.count("lastRenderedPageBreak"),
        "lang_tags": dict(langs.most_common(4)),
        "comments": "word/comments.xml" in names,
        "tracked_changes": body.count("<w:ins ") + body.count("<w:del "),
    }


def main():
    data_dir = Path(sys.argv[1] if len(sys.argv) > 1 else "data")
    out = Path(sys.argv[2] if len(sys.argv) > 2 else "analysis/inventory.json")
    results = {}
    for path in sorted(data_dir.rglob("*")):
        if not path.is_file() or path.name.startswith("."):
            continue
        rel = str(path.relative_to(data_dir))
        info = {"size_kb": round(path.stat().st_size / 1024, 1)}
        try:
            if path.suffix.lower() == ".pdf":
                info |= inspect_pdf(path)
            elif path.suffix.lower() == ".docx":
                info |= inspect_docx(path)
            else:
                info["type"] = path.suffix.lstrip(".") or "other"
        except Exception as exc:
            info["error"] = repr(exc)
        results[rel] = info
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(results, indent=1, ensure_ascii=False))
    print(f"Inspected {len(results)} files -> {out}")


if __name__ == "__main__":
    main()
