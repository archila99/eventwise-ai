"""Lightweight fixtures for indexing/retrieval tests (no real embedding model)."""

from __future__ import annotations

import hashlib
from pathlib import Path

import docx
import numpy as np

from eventwise.config import IndexConfig
from eventwise.indexing.embeddings import normalise_rows
from eventwise.indexing.index import LoadedIndex
from eventwise.indexing.units import build_units
from eventwise.models import CaptionMatch, Element, ExtractedDocument, FigureData, NormalisedCell, TableData
from eventwise.retrieval.tokenize import tokenize


class FakeEmbedder:
    """Deterministic hashed bag-of-words vectors: texts sharing words are similar."""

    def __init__(self, dim: int = 64, name: str = "fake:bow-64"):
        self.dim = dim
        self.name = name
        self.calls = 0
        self.texts_encoded = 0

    def encode(self, texts, batch_size: int = 64) -> np.ndarray:
        self.calls += 1
        self.texts_encoded += len(texts)
        out = np.zeros((len(texts), self.dim), dtype=np.float32)
        for i, text in enumerate(texts):
            for token in tokenize(text):
                h = int(hashlib.md5(token.encode()).hexdigest(), 16)
                out[i, h % self.dim] += 1.0 if (h >> 8) & 1 else -1.0
        out[np.linalg.norm(out, axis=1) == 0, 0] = 1.0
        return normalise_rows(out)


def table(rows: list[list[str]], header_rows: int = 1) -> TableData:
    return TableData(n_rows=len(rows), n_cols=len(rows[0]), raw_cells=[list(r) for r in rows],
                     normalised_cells=[[NormalisedCell(c) for c in r] for r in rows], header_rows=header_rows)


def make_doc(doc_id: str, name: str, elements: list[dict], title: str | None = None) -> dict:
    """A document dict in the Stage 1 ``document.json`` schema."""
    doc = ExtractedDocument(doc_id=doc_id, doc_name=name, source_path=f"/tmp/{name}", file_type="pdf",
                            metadata={"title": title} if title else {})
    short = doc_id[:10]
    for order, spec in enumerate(elements):
        spec = dict(spec)
        kind = spec.pop("kind")
        page = spec.pop("page", 1)
        eid = spec.pop("element_id", f"{short}:p{page}:{kind[:3]}{order}")
        doc.elements.append(Element(element_id=eid, kind=kind, order=order, page=page,
                                    text=spec.pop("text", ""), **spec))
    return doc.to_dict()


def figure(image: str = "images/p1_fig1.png", labels: str | None = None) -> FigureData:
    return FigureData(source_type="raster", image_path=image, text_labels=labels)


def matched() -> CaptionMatch:
    return CaptionMatch(status="matched", confidence="high", method="geometry")


def loaded_index(docs: list[dict], embedder: FakeEmbedder, config: IndexConfig | None = None,
                 extracted_dir: Path = Path("/tmp/extracted")) -> LoadedIndex:
    """An in-memory LoadedIndex built from document dicts (bypasses disk)."""
    units = [u for d in docs for u in build_units(d, config)]
    vectors = embedder.encode([u.search_text for u in units]) if units else np.zeros((0, embedder.dim), np.float32)
    manifest = {"embedder": embedder.name, "dim": embedder.dim,
                "documents": {d["doc_id"]: {"doc_name": d["doc_name"], "source_path": d["source_path"],
                                            "n_units": sum(u.doc_id == d["doc_id"] for u in units)} for d in docs}}
    return LoadedIndex(units=units, vectors=vectors, manifest=manifest, index_dir=Path("/tmp/index"),
                       extracted_dir=extracted_dir)


def write_docx(path: Path, heading: str, paragraphs: list[str]) -> Path:
    d = docx.Document()
    d.add_heading(heading, level=1)
    for p in paragraphs:
        d.add_paragraph(p)
    d.save(str(path))
    return path
