"""Document-agnostic behaviour on synthetic and broken inputs."""

import docx
import pymupdf
import pytest

from eventwise.extraction.extract import extract_document, parse_page_spec


@pytest.mark.parametrize("name", ["broken.pdf", "broken.docx"])
def test_corrupt_files_fail_gracefully(tmp_path, name):
    path = tmp_path / name
    path.write_bytes(b"this is not a real document")
    doc = extract_document(path)
    assert doc.status == "failed"
    assert doc.issues[0].code == "parse_failed" and doc.elements == []


def test_unsupported_and_missing_files(tmp_path):
    txt = tmp_path / "notes.txt"
    txt.write_text("hello")
    assert extract_document(txt).issues[0].code == "unsupported_file_type"
    assert extract_document(tmp_path / "missing.pdf").issues[0].code == "unreadable_file"


def _synthetic_pdf(path):
    pdf = pymupdf.open()
    for n in range(1, 4):
        page = pdf.new_page()
        page.insert_text((72, 30), "Acme Ltd - Quarterly Report", fontsize=9)
        page.insert_text((72, 100), f"{n}. Section heading", fontsize=18)
        page.insert_text((72, 140), "Ordinary body text about the quarter.", fontsize=11)
        page.insert_text((72, 180), "Ordinary body text continues here.", fontsize=11)
        page.insert_text((300, 820), str(n), fontsize=9)
    page = pdf[0]
    y = 220
    for k in range(4):  # ruled 3x2 grid
        page.draw_line((72, y + k * 20), (312, y + k * 20))
    for x in (72, 192, 312):
        page.draw_line((x, y), (x, y + 60))
    for row in range(3):
        for col in range(2):
            page.insert_text((76 + col * 120, y + 14 + row * 20), ["Item", "Cost"][col] if row == 0 else
                             [f"Item {row}", f"£{row},000"][col], fontsize=10)
    page.insert_text((72, 300), "Table 1: Example costs.", fontsize=10)
    pdf[1].insert_text((72, 600), "Figure 9: A caption with no figure.", fontsize=10)
    pdf.save(path)


def test_synthetic_pdf(tmp_path):
    path = tmp_path / "synthetic.pdf"
    _synthetic_pdf(path)
    doc = extract_document(path)
    assert doc.status == "ok"
    table = doc.by_kind("table")[0]
    assert table.label == "Table 1"
    assert table.table.normalised_cells[1][1].value == 1000.0
    assert [h.text for h in doc.by_kind("heading")] == ["1. Section heading", "2. Section heading", "3. Section heading"]
    assert {p["label"] for p in doc.pages} == {"1", "2", "3"}
    assert all("Acme Ltd" not in e.text for e in doc.elements)
    unmatched = [i for i in doc.issues if i.code == "caption_unmatched"]
    assert len(unmatched) == 1 and unmatched[0].location["page"] == 2
    assert any(e.text.startswith("Figure 9") for e in doc.elements), "unmatched captions stay as text"


def test_synthetic_docx(tmp_path):
    d = docx.Document()
    d.add_heading("Overview", level=1)
    d.add_paragraph("Body text.")
    table = d.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged header"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Spans two rows"
    table.cell(1, 2).text = "(£1,200)"
    d.add_paragraph("Table 1: Synthetic table.")
    d.add_heading("Details", level=2)
    d.add_paragraph("Table 7 is referenced here but this is not a caption.")
    path = tmp_path / "synthetic.docx"
    d.save(path)

    doc = extract_document(path)
    assert doc.status == "ok"
    tbl = doc.by_kind("table")[0]
    assert tbl.label == "Table 1" and tbl.section_path == ["Overview"]
    merges = {(m.row, m.col): (m.row_span, m.col_span) for m in tbl.table.merges}
    assert merges == {(0, 0): (1, 2), (1, 0): (2, 1)}
    assert tbl.table.raw_cells[0][1] is None and tbl.table.raw_cells[2][0] is None
    assert tbl.table.normalised_cells[1][2].value == -1200.0
    last = doc.elements[-1]
    assert last.kind == "text" and last.section_path == ["Overview", "Details"]
    assert not [i for i in doc.issues if i.code == "caption_unmatched"]


def test_invalid_pages_reported(tmp_path):
    path = tmp_path / "synthetic.pdf"
    _synthetic_pdf(path)
    doc = extract_document(path, pages=[2, 99])
    assert doc.pages_processed == [2]
    assert any(i.code == "invalid_pages" for i in doc.issues)


def test_parse_page_spec():
    assert parse_page_spec("1-3,7, 9") == [1, 2, 3, 7, 9]
