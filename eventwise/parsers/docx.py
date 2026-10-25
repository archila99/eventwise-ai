"""DOCX extraction with python-docx plus direct WordprocessingML access.

The body is walked in document order so headings, paragraphs, tables and
captions keep their sequence. Content python-docx does not expose (text
boxes, legacy form checkboxes, symbols, merged-cell spans, images, native
charts) is read from the XML directly. DOCX has no fixed pages, so locations
are block/paragraph/table indices plus the section path.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from lxml import etree

from ..captions import Caption, match_by_adjacency, parse_caption
from ..models import (
    CaptionMatch,
    CellMerge,
    Element,
    ExtractedDocument,
    FigureData,
    TableData,
)
from ..normalise import clean_text, normalise_cell

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "v": "urn:schemas-microsoft-com:vml",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
}


def _q(tag: str) -> str:
    prefix, local = tag.split(":")
    return f"{{{NS[prefix]}}}{local}"


W_P, W_TBL, W_SDT, W_SDTCONTENT, W_SECTPR = map(_q, ("w:p", "w:tbl", "w:sdt", "w:sdtContent", "w:sectPr"))
W_T, W_TAB, W_BR, W_CR, W_SYM = map(_q, ("w:t", "w:tab", "w:br", "w:cr", "w:sym"))
W_DELTEXT, W_INSTR, W_TXBX, W_CHECKBOX = map(_q, ("w:delText", "w:instrText", "w:txbxContent", "w:checkBox"))
W14_CHECKBOX, MC_FALLBACK, A_BLIP, V_IMAGEDATA, C_CHART = map(
    _q, ("w14:checkbox", "mc:Fallback", "a:blip", "v:imagedata", "c:chart"))
W_TR, W_TC, W_VAL = _q("w:tr"), _q("w:tc"), _q("w:val")
W_NOBREAKHYPHEN, W_INS, W_DEL = _q("w:noBreakHyphen"), _q("w:ins"), _q("w:del")
R_EMBED, R_ID = _q("r:embed"), _q("r:id")

UNCHECKED, CHECKED = "\u2610", "\u2612"
NUMBERED_HEADING = re.compile(r"^(\d+(?:\.\d+)*)\.?\s*[^\W\d_]")
EMU_PER_PX = 9525


@dataclass
class _Content:
    """Everything found inside one paragraph (or cell paragraph)."""

    parts: list[str] = field(default_factory=list)
    checkboxes: list[dict] = field(default_factory=list)
    symbols: list[str] = field(default_factory=list)
    images: list[str] = field(default_factory=list)  # relationship ids
    charts: list[str] = field(default_factory=list)
    textboxes: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "".join(self.parts)


def _on_off(el, default: bool = False) -> bool:
    """WordprocessingML on/off property: a bare <w:b/> means on."""
    if el is None:
        return default
    val = el.get(W_VAL) or el.get(_q("w14:val"))
    return val is None or val.lower() not in ("0", "false", "off")


def _walk(el, acc: _Content) -> None:
    for child in el:
        tag = child.tag
        if not isinstance(tag, str) or tag == MC_FALLBACK:  # VML fallback duplicates DrawingML
            continue
        if tag == W_TXBX:
            text = "\n".join(t for t in (_paragraph_content(p).text.strip() for p in child.iter(W_P)) if t)
            if text:
                acc.textboxes.append(text)
        elif tag == W_T:
            acc.parts.append(child.text or "")
        elif tag == W_TAB:
            acc.parts.append("\t")
        elif tag in (W_BR, W_CR):
            acc.parts.append("\n")
        elif tag == W_NOBREAKHYPHEN:
            acc.parts.append("-")
        elif tag == W_SYM:
            font, char = child.get(_q("w:font")), child.get(_q("w:char"))
            acc.symbols.append(f"{font}:{char}")
            acc.parts.append(f"[symbol {font} {char}]")
        elif tag == W_CHECKBOX:  # legacy form field: <w:ffData><w:checkBox>
            checked = _on_off(child.find(_q("w:checked")), _on_off(child.find(_q("w:default"))))
            acc.checkboxes.append({"type": "legacy_form_checkbox", "checked": checked})
            acc.parts.append(CHECKED if checked else UNCHECKED)
        elif tag == W14_CHECKBOX:  # content control; its glyph is already in the text
            checked = _on_off(child.find(_q("w14:checked")))
            acc.checkboxes.append({"type": "content_control_checkbox", "checked": checked})
        elif tag == A_BLIP and child.get(R_EMBED):
            acc.images.append(child.get(R_EMBED))
        elif tag == V_IMAGEDATA and child.get(R_ID):
            acc.images.append(child.get(R_ID))
        elif tag == C_CHART and child.get(R_ID):
            acc.charts.append(child.get(R_ID))
        elif tag in (W_DELTEXT, W_INSTR):
            continue
        else:
            _walk(child, acc)


def _paragraph_content(p) -> _Content:
    acc = _Content()
    _walk(p, acc)
    return acc


def _style_name(p, styles: dict[str, str]) -> str:
    style = p.find(f"{_q('w:pPr')}/{_q('w:pStyle')}")
    style_id = style.get(W_VAL) if style is not None else None
    return styles.get(style_id, style_id or "Normal")


def _heading_level(p, style: str) -> int | None:
    lower = style.lower()
    if lower == "title":
        return 0
    match = re.match(r"heading\s*(\d)", lower)
    if match:
        return int(match.group(1))
    outline = p.find(f"{_q('w:pPr')}/{_q('w:outlineLvl')}")
    if outline is not None and outline.get(W_VAL, "").isdigit() and int(outline.get(W_VAL)) < 9:
        return int(outline.get(W_VAL)) + 1
    return None


def _all_runs_bold(p) -> bool:
    runs = [r for r in p.iter(_q("w:r")) if "".join(t.text or "" for t in r.iter(W_T)).strip()]
    return bool(runs) and all(_on_off(r.find(f"{_q('w:rPr')}/{_q('w:b')}")) for r in runs)


@dataclass
class _Pending:
    """An element before captions and section paths are resolved."""

    kind: str  # heading | text | table | figure | caption
    element: Element
    caption: Caption | None = None


class _DocxParser:
    def __init__(self, path: Path, doc: ExtractedDocument, out_dir: Path | None):
        self.path, self.doc, self.out_dir = path, doc, out_dir
        self.document: DocxDocument = docx.Document(str(path))
        self.styles = {s.style_id: s.name for s in self.document.styles if s.style_id}
        self.short = doc.doc_id[:10]
        self.items: list[_Pending] = []
        self.paragraph_index = 0
        self.table_index = 0
        self.uses_heading_styles = any(
            _heading_level(p, _style_name(p, self.styles)) not in (None, 0)
            for p in self.document.element.body.iter(W_P)
        )

    # ---------------------------------------------------------- body walk

    def run(self) -> ExtractedDocument:
        self._metadata()
        for block_index, child in enumerate(self._blocks(self.document.element.body)):
            try:
                if child.tag == W_P:
                    self._paragraph(child, block_index)
                elif child.tag == W_TBL:
                    self._table(child, block_index)
                else:
                    self.doc.add_issue("unsupported_block", f"Body element <{etree.QName(child).localname}> skipped",
                                       block=block_index)
            except Exception as exc:
                self.doc.add_issue("block_failed", f"Block could not be processed: {exc}", severity="error",
                                   block=block_index)
        self._notes()
        self._resolve()
        return self.doc

    def _blocks(self, parent):
        for child in parent:
            if child.tag == W_SECTPR or not isinstance(child.tag, str):
                continue
            if child.tag == W_SDT:  # block-level content control: unwrap
                content = child.find(W_SDTCONTENT)
                if content is not None:
                    yield from self._blocks(content)
            elif etree.QName(child).localname in ("customXml", "smartTag"):
                yield from self._blocks(child)
            else:
                yield child

    def _paragraph(self, p, block_index: int) -> None:
        self.paragraph_index += 1
        content = _paragraph_content(p)
        style = _style_name(p, self.styles)
        raw = content.text
        text = clean_text(raw)
        location = {"block": block_index, "paragraph": self.paragraph_index, "style": style}
        for cb in content.checkboxes:
            self.doc.form_fields.append({**cb, **location})
        if content.symbols:
            self.doc.add_issue("unmapped_symbol", f"Symbol font characters kept as placeholders: {content.symbols}",
                               severity="info", block=block_index)

        if text:
            level = _heading_level(p, style)
            heading_source = "style" if level is not None else None
            if level is None and not self.uses_heading_styles and len(text) <= 80 \
                    and NUMBERED_HEADING.match(text) and _all_runs_bold(p):
                level = NUMBERED_HEADING.match(text).group(1).count(".") + 1
                heading_source = "numbered_bold"
            caption = parse_caption(text) if level is None else None
            if caption is None and style.lower() == "caption":
                caption = Caption(kind="figure", label="", text=text)
            kind = "heading" if level is not None else ("caption" if caption else "text")
            element = Element(
                element_id=f"{self.short}:b{block_index}:{'h' if kind == 'heading' else 'x'}",
                kind="heading" if kind == "heading" else "text", order=0,
                text=text, raw_text=raw, heading_level=level,
                location={**location, **({"heading_source": heading_source} if heading_source else {})},
            )
            self.items.append(_Pending(kind, element, caption))

        for n, box_text in enumerate(content.textboxes, 1):
            self.items.append(_Pending("text", Element(
                element_id=f"{self.short}:b{block_index}:textbox{n}", kind="text", order=0,
                text=clean_text(box_text), raw_text=box_text,
                location={**location, "container": "textbox"},
            )))
        for n, rid in enumerate(content.images, 1):
            self._image(rid, block_index, n, location)
        for n, rid in enumerate(content.charts, 1):
            self._chart(rid, block_index, n, location)

    # ---------------------------------------------------------- tables

    def _table(self, tbl, block_index: int) -> None:
        self.table_index += 1
        location = {"block": block_index, "table": self.table_index}
        data, cell_images = self._table_data(tbl, location)
        element = Element(
            element_id=f"{self.short}:b{block_index}:tab", kind="table", order=0,
            text="\n".join(" | ".join(c.text if c else "" for c in row) for row in data.normalised_cells),
            location=location, table=data,
        )
        self.items.append(_Pending("table", element))
        for n, (rid, cell) in enumerate(cell_images, 1):
            self._image(rid, block_index, n, {**location, "cell": cell}, suffix="cellfig")

    def _table_data(self, tbl, location: dict) -> tuple[TableData, list]:
        origins: dict[tuple[int, int], dict] = {}
        covered: dict[tuple[int, int], tuple[int, int]] = {}
        open_vmerge: dict[int, tuple[int, int]] = {}
        cell_images = []
        rows = tbl.findall(W_TR)
        n_cols = 0
        for r, tr in enumerate(rows):
            grid_before = tr.find(f"{_q('w:trPr')}/{_q('w:gridBefore')}")
            c = int(grid_before.get(W_VAL, 0)) if grid_before is not None else 0
            for tc in tr.iter(W_TC):
                if tc.getparent() is not tr and tc.getparent().getparent() is not tr:
                    continue  # skip cells of nested tables
                tc_pr = tc.find(_q("w:tcPr"))
                span_el = tc_pr.find(_q("w:gridSpan")) if tc_pr is not None else None
                span = int(span_el.get(W_VAL, 1)) if span_el is not None else 1
                vmerge_el = tc_pr.find(_q("w:vMerge")) if tc_pr is not None else None
                vmerge = None if vmerge_el is None else (vmerge_el.get(W_VAL) or "continue")
                text, images = self._cell_text(tc, {**location, "row": r, "col": c})
                cell_images.extend((rid, [r, c]) for rid in images)

                if vmerge == "continue" and c in open_vmerge:
                    origin = open_vmerge[c]
                    for k in range(span):
                        covered[(r, c + k)] = origin
                    if text.strip():
                        self.doc.add_issue("merged_cell_text", "Text found in a vertically merged continuation cell",
                                           severity="info", **location, row=r, col=c)
                else:
                    origins[(r, c)] = {"raw": text, "row_span": 1, "col_span": span}
                    for k in range(1, span):
                        covered[(r, c + k)] = (r, c)
                    if vmerge == "restart":
                        open_vmerge[c] = (r, c)
                    else:
                        open_vmerge.pop(c, None)
                c += span
            n_cols = max(n_cols, c)

        for (r, c), (r0, c0) in covered.items():
            o = origins[(r0, c0)]
            o["row_span"] = max(o["row_span"], r - r0 + 1)
            o["col_span"] = max(o["col_span"], c - c0 + 1)

        raw = [[origins[(r, c)]["raw"] if (r, c) in origins else None for c in range(n_cols)] for r in range(len(rows))]
        missing = sum(1 for r in range(len(rows)) for c in range(n_cols) if (r, c) not in origins and (r, c) not in covered)
        if missing:
            self.doc.add_issue("ragged_table", f"{missing} grid position(s) have no cell (ragged rows)",
                               severity="info", **location)
        merges = [CellMerge(r, c, o["row_span"], o["col_span"]) for (r, c), o in sorted(origins.items())
                  if o["row_span"] > 1 or o["col_span"] > 1]
        data = TableData(
            n_rows=len(rows), n_cols=n_cols, raw_cells=raw,
            normalised_cells=[[normalise_cell(v) for v in row] for row in raw],
            merges=merges, header_rows=1, method="docx_xml",
        )
        return data, cell_images

    def _cell_text(self, tc, location: dict) -> tuple[str, list[str]]:
        parts, images = [], []
        for child in tc:
            if child.tag == W_P:
                content = _paragraph_content(child)
                parts.append(content.text)
                parts.extend(content.textboxes)
                images.extend(content.images)
                for cb in content.checkboxes:
                    self.doc.form_fields.append({**cb, **location})
                if content.symbols:
                    self.doc.add_issue("unmapped_symbol", f"Symbol font characters kept as placeholders: {content.symbols}",
                                       severity="info", **location)
                if content.textboxes:
                    self.doc.add_issue("textbox_in_table_cell", "Text box content appended to its table cell text",
                                       severity="info", **location)
            elif child.tag == W_TBL:
                nested = "\n".join(" | ".join(clean_text(self._cell_text(c, location)[0]) for c in tr.findall(W_TC))
                                   for tr in child.findall(W_TR))
                parts.append(nested)
                self.doc.add_issue("nested_table_flattened", "Nested table flattened into its parent cell text",
                                   **location)
            elif child.tag == W_SDT:
                text, imgs = self._cell_text(child.find(W_SDTCONTENT) if child.find(W_SDTCONTENT) is not None
                                             else child, location)
                parts.append(text)
                images.extend(imgs)
        return "\n".join(parts), images

    # ---------------------------------------------------------- figures

    def _image(self, rid: str, block_index: int, n: int, location: dict, suffix: str = "fig") -> None:
        element_id = f"{self.short}:b{block_index}:{suffix}{n}"
        figure = FigureData("embedded_image")
        try:
            part = self.document.part.related_parts[rid]
            blob = part.blob
            ext = Path(str(part.partname)).suffix.lstrip(".").lower() or "bin"
            figure.image_format = ext
            if ext in ("png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff"):
                try:
                    import pymupdf
                    pix = pymupdf.Pixmap(blob)
                    figure.width_px, figure.height_px = pix.width, pix.height
                except Exception:
                    pass
            else:
                self.doc.add_issue("image_format_not_previewable", f"Embedded image in {ext} format saved but not rendered",
                                   severity="info", element_id=element_id)
            if self.out_dir is not None:
                image_dir = self.out_dir / "images"
                image_dir.mkdir(parents=True, exist_ok=True)
                name = f"{element_id.split(':', 1)[1].replace(':', '_')}.{ext}"
                (image_dir / name).write_bytes(blob)
                figure.image_path = f"images/{name}"
        except KeyError:
            self.doc.add_issue("image_missing", f"Image relationship {rid} not found (linked or external image?)",
                               element_id=element_id)
        self.items.append(_Pending("figure", Element(
            element_id=element_id, kind="figure", order=0, text="",
            location={**location, "relationship_id": rid}, figure=figure,
        )))

    def _chart(self, rid: str, block_index: int, n: int, location: dict) -> None:
        element_id = f"{self.short}:b{block_index}:chart{n}"
        figure = FigureData("native_chart")
        try:
            root = etree.fromstring(self.document.part.related_parts[rid].blob)
            figure.chart_data = _chart_series(root)
            title = " ".join(t.text or "" for t in root.iterfind(".//c:title//a:t", NS)).strip()
            text = clean_text(" ".join([title] + [f"{s['name']}: " + ", ".join(
                f"{c}={v}" for c, v in zip(s["categories"], s["values"])) for s in figure.chart_data]))
        except Exception as exc:
            text = ""
            self.doc.add_issue("chart_data_unreadable", f"Native chart could not be read: {exc}", element_id=element_id)
        self.items.append(_Pending("figure", Element(
            element_id=element_id, kind="figure", order=0, text=text,
            location={**location, "relationship_id": rid}, figure=figure,
        )))

    # ---------------------------------------------------------- metadata & notes

    def _metadata(self) -> None:
        props = self.document.core_properties
        self.doc.metadata = {k: str(v) for k, v in {
            "title": props.title, "author": props.author, "created": props.created,
            "modified": props.modified}.items() if v}
        seen = set()
        for section in self.document.sections:
            for kind in ("header", "footer", "first_page_header", "first_page_footer",
                         "even_page_header", "even_page_footer"):
                part = getattr(section, kind)
                if part.is_linked_to_previous:
                    continue
                text = clean_text(" ".join(_paragraph_content(p).text for p in part._element.iter(W_P)))
                if text and text not in seen:
                    seen.add(text)
                    self.doc.furniture.append({"text": text, "kind": kind})
        body = self.document.element.body
        if body.find(f".//{W_INS}") is not None or body.find(f".//{W_DEL}") is not None:
            self.doc.add_issue("tracked_changes", "Tracked changes present: insertions kept, deletions dropped",
                               severity="info")

    def _notes(self) -> None:
        """Footnotes/endnotes become trailing text elements; comments are reported only."""
        for rel in self.document.part.rels.values():
            kind = rel.reltype.rsplit("/", 1)[-1]
            if kind not in ("footnotes", "endnotes", "comments") or rel.is_external:
                continue
            root = etree.fromstring(rel.target_part.blob)
            tag = {"footnotes": "w:footnote", "endnotes": "w:endnote", "comments": "w:comment"}[kind]
            notes = [n for n in root.iterfind(tag, NS) if n.get(_q("w:type")) in (None, "normal")]
            if kind == "comments":
                if notes:
                    self.doc.add_issue("comments_not_extracted", f"{len(notes)} reviewer comment(s) not extracted",
                                       severity="info")
                continue
            for note in notes:
                text = clean_text(" ".join(_paragraph_content(p).text for p in note.iter(W_P)))
                if text:
                    self.items.append(_Pending("text", Element(
                        element_id=f"{self.short}:{kind}{note.get(_q('w:id'))}", kind="text", order=0,
                        text=text, raw_text=text, location={"container": kind[:-1], "note_id": note.get(_q("w:id"))},
                    )))

    # ---------------------------------------------------------- captions & sections

    def _resolve(self) -> None:
        sequence = [(p.kind, p.caption) for p in self.items]
        assignments, ambiguous = match_by_adjacency(sequence)
        consumed = set()
        for a in assignments:
            cap_item, obj = self.items[a.caption_index], self.items[a.object_index].element
            caption = cap_item.caption
            obj.caption, obj.label = caption.text, caption.label or None
            obj.caption_match = CaptionMatch("matched", a.confidence, a.method, None,
                                             {k: v for k, v in cap_item.element.location.items()
                                              if k in ("block", "paragraph")})
            if obj.kind == "figure":
                obj.text = clean_text(f"{caption.text} {obj.text}")
            consumed.add(a.caption_index)
            if a.confidence == "low":
                self.doc.add_issue("caption_low_confidence",
                                   f"'{caption.label}' attached to a {obj.kind} although the label suggests a {caption.kind}",
                                   element_id=obj.element_id)
        for i in ambiguous:
            self.doc.add_issue("caption_ambiguous", "Caption sits between two candidate objects; attached to the preceding one",
                               element_id=self.items[i].element.element_id)
        for i, item in enumerate(self.items):
            if item.kind == "caption" and i not in consumed:
                self.doc.add_issue("caption_unmatched", f"No adjacent {item.caption.kind} for caption '{item.caption.label}'",
                                   element_id=item.element.element_id)

        stack: list[tuple[int, str]] = []
        order = 0
        for i, item in enumerate(self.items):
            if i in consumed:
                continue
            element = item.element
            if element.kind == "heading":
                while stack and stack[-1][0] >= element.heading_level:
                    stack.pop()
                element.section_path = [t for _, t in stack]
                stack.append((element.heading_level, element.text))
            else:
                element.section_path = [t for _, t in stack]
            if element.kind in ("table", "figure") and element.caption_match is None:
                element.caption_match = CaptionMatch("unmatched")
            element.order = order
            order += 1
            self.doc.elements.append(element)

        uncaptioned = [e.element_id for e in self.doc.elements if e.kind in ("table", "figure") and not e.caption]
        if uncaptioned:
            self.doc.add_issue("objects_without_caption", f"{len(uncaptioned)} table(s)/figure(s) have no caption",
                               severity="info", element_ids=uncaptioned)


def _chart_series(root) -> list[dict]:
    series = []
    for ser in root.iterfind(".//c:ser", NS):
        name = " ".join(v.text or "" for v in ser.iterfind("c:tx//c:v", NS)).strip()
        categories = [v.text for v in ser.iterfind("c:cat//c:pt/c:v", NS)]
        values = [v.text for v in ser.iterfind("c:val//c:pt/c:v", NS)]
        series.append({"name": name, "categories": categories, "values": values})
    return series


def parse_docx(path: Path, doc: ExtractedDocument, out_dir: Path | None = None) -> ExtractedDocument:
    return _DocxParser(path, doc, out_dir).run()
